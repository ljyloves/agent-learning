"""Teacher-answer and answer-sheet Word documents for an assembled paper."""

from __future__ import annotations

import asyncio
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Mapping

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    KnowledgePointModel,
    QuestionAnalysisModel,
    QuestionKnowledgePointModel,
)
from app.modules.paper_agent.schemas.analysis import DifficultyEstimationLLMOutput
from app.modules.paper_agent.schemas.question import Question, QuestionType
from app.modules.paper_agent.schemas.student_paper import (
    AnswerSheetWordRequest,
    TeacherAnswerWordRequest,
)
from app.modules.paper_agent.services.resource_storage import (
    get_stored_resource_file,
    recover_question,
)
from app.modules.paper_agent.services.student_paper import (
    CHINESE_SECTION_NUMBERS,
    DOCX_MIME_TYPE,
    SECTION_TITLES,
    StudentPaperExportDataError,
    StudentPaperTemplateError,
    StudentPaperWordFile,
    _add_images,
    _add_options,
    _clear_template_body,
    _configure_a4_section,
    _ensure_styles,
    _prevent_row_split,
    _question_images,
    _remove_table_borders,
    _request_field_updates,
    _set_cell_margins,
    create_student_paper_template,
    ordered_assembled_questions,
    protected_answer_texts,
    safe_docx_filename,
)


TEACHER_TEMPLATE_FILENAME = "teacher_answer.docx"
ANSWER_SHEET_TEMPLATE_FILENAME = "answer_sheet.docx"
OBJECTIVE_TYPES = {
    QuestionType.SINGLE_CHOICE,
    QuestionType.MULTIPLE_CHOICE,
    QuestionType.TRUE_FALSE,
}


@dataclass(frozen=True, slots=True)
class TeacherQuestionMetadata:
    knowledge_points: tuple[tuple[str, str], ...]
    difficulty_level: int
    estimated_correct_rate: float


def create_teacher_answer_template() -> DocumentObject:
    document = create_student_paper_template()
    properties = document.core_properties
    properties.title = "FlowGate 高中生物教师答案解析模板"
    properties.subject = "A4 teacher answer and explanation template"
    properties.keywords = "FlowGate,高中生物,教师卷,答案,解析,A4"
    return document


def create_answer_sheet_template() -> DocumentObject:
    document = create_student_paper_template()
    properties = document.core_properties
    properties.title = "FlowGate 高中生物答题卡模板"
    properties.subject = "A4 answer sheet template"
    properties.keywords = "FlowGate,高中生物,答题卡,A4"
    return document


def _open_template(path: Path, label: str) -> DocumentObject:
    if not path.is_file():
        raise StudentPaperTemplateError(f"{label} template is missing: {path}")
    try:
        document = Document(path)
    except Exception as exc:
        raise StudentPaperTemplateError(f"{label} template cannot be opened") from exc
    _clear_template_body(document)
    _ensure_styles(document)
    for section in document.sections:
        _configure_a4_section(section)
    _request_field_updates(document)
    return document


def _section_number(index: int) -> str:
    return CHINESE_SECTION_NUMBERS[index] if index < 10 else str(index + 1)


def _answer_text(answer: str | list[str] | None) -> str:
    if isinstance(answer, list):
        return "；".join(answer)
    return answer or ""


def _validate_teacher_question(question: Question, metadata: TeacherQuestionMetadata) -> None:
    answer_targets = question.subquestions or [question]
    for index, target in enumerate(answer_targets, start=1):
        location = f" subquestion {index}" if question.subquestions else ""
        if not _answer_text(target.answer):
            raise StudentPaperExportDataError(
                f"teacher answer is missing for question {question.id}{location}"
            )
        if not target.explanation:
            raise StudentPaperExportDataError(
                f"teacher explanation is missing for question {question.id}{location}"
            )
    if not metadata.knowledge_points:
        raise StudentPaperExportDataError(
            f"knowledge points are missing for question {question.id}"
        )


def _add_teacher_details(
    document: DocumentObject,
    question: Question,
    metadata: TeacherQuestionMetadata,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    row = table.rows[0]
    _prevent_row_split(row)
    cell = row.cells[0]
    _set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
    paragraph = cell.paragraphs[0]
    for label, value in (
        (
            "答案",
            _answer_text(question.answer)
            if not question.subquestions
            else "；".join(
                f"（{index}）{_answer_text(child.answer)}"
                for index, child in enumerate(question.subquestions, start=1)
            ),
        ),
        (
            "解析",
            question.explanation
            if not question.subquestions
            else "\n".join(
                f"（{index}）{child.explanation}"
                for index, child in enumerate(question.subquestions, start=1)
            ),
        ),
        (
            "知识点",
            "；".join(f"{code} {name}" for code, name in metadata.knowledge_points),
        ),
        (
            "难度",
            f"{metadata.difficulty_level} 级（预计正确率 "
            f"{metadata.estimated_correct_rate:.0%}）",
        ),
    ):
        heading = paragraph.add_run(f"{label}：")
        heading.bold = True
        paragraph.add_run(value)
        paragraph = cell.add_paragraph()
    cell._tc.remove(paragraph._p)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def render_teacher_answer_document(
    request: TeacherAnswerWordRequest,
    *,
    questions: Mapping[str, Question],
    metadata: Mapping[str, TeacherQuestionMetadata],
    image_data: Mapping[str, bytes],
    template_path: Path,
) -> bytes:
    document = _open_template(template_path, "teacher answer")
    title = document.add_paragraph(request.title, style="PaperTitle")
    title.paragraph_format.keep_with_next = True
    subtitle = document.add_paragraph(
        request.subtitle or f"{request.paper.module_code} 教师用卷",
        style="PaperSubtitle",
    )
    subtitle.paragraph_format.keep_with_next = True

    number_by_id = {
        assembled.question_id: number
        for number, assembled in ordered_assembled_questions(request.paper)
    }
    for section_position, assembled_section in enumerate(request.paper.sections):
        if section_position > 0 and request.start_each_section_on_new_page:
            document.add_page_break()
        heading = document.add_paragraph(
            f"{_section_number(section_position)}、"
            f"{SECTION_TITLES[assembled_section.question_type]}",
            style="PaperSection",
        )
        heading.paragraph_format.keep_with_next = True
        for assembled in assembled_section.questions:
            question = questions.get(assembled.question_id)
            question_metadata = metadata.get(assembled.question_id)
            if question is None or question_metadata is None:
                raise StudentPaperExportDataError(
                    f"teacher export data is missing: {assembled.question_id}"
                )
            if question.question_type != assembled.question_type:
                raise StudentPaperExportDataError(
                    f"question type changed after assembly: {assembled.question_id}"
                )
            _validate_teacher_question(question, question_metadata)
            block = document.add_table(rows=1, cols=1)
            block.alignment = WD_TABLE_ALIGNMENT.CENTER
            block.autofit = True
            _remove_table_borders(block)
            row = block.rows[0]
            _prevent_row_split(row)
            cell = row.cells[0]
            _set_cell_margins(cell, top=40, bottom=80, left=0, right=0)
            prompt = cell.paragraphs[0]
            prompt.paragraph_format.keep_with_next = True
            number = prompt.add_run(f"{number_by_id[assembled.question_id]}. ")
            number.bold = True
            prompt.add_run(question.stem)
            prompt.add_run(f"（{assembled.score} 分）").font.size = Pt(9)
            _add_images(cell, question.images, image_data)
            _add_options(
                cell,
                question,
                image_data,
                keep_with_next=True,
            )
            for index, child in enumerate(question.subquestions, start=1):
                child_prompt = cell.add_paragraph(f"（{index}）{child.stem}")
                child_prompt.paragraph_format.keep_with_next = bool(child.images)
                _add_images(cell, child.images, image_data)
                _add_options(cell, child, image_data, keep_with_next=True)
            _add_teacher_details(cell, question, question_metadata)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _add_identity(document: DocumentObject, request: AnswerSheetWordRequest) -> None:
    title = document.add_paragraph(request.title, style="PaperTitle")
    title.paragraph_format.keep_with_next = True
    subtitle = document.add_paragraph(
        request.subtitle or f"{request.paper.module_code} 答题卡",
        style="PaperSubtitle",
    )
    subtitle.paragraph_format.keep_with_next = True
    table = document.add_table(rows=2, cols=2)
    _remove_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = (
        f"学校：{request.school_name or '____________'}",
        f"年级：{request.grade_name}",
        "姓名：____________",
        "考号：____________",
    )
    for cell, value in zip(
        (cell for row in table.rows for cell in row.cells),
        values,
        strict=True,
    ):
        _set_cell_margins(cell, top=30, bottom=30, left=60, right=60)
        cell.paragraphs[0].add_run(value)


def _add_objective_answer_rows(
    document: DocumentObject,
    entries: list[tuple[int, Question]],
) -> None:
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for offset in range(0, len(entries), 2):
        row = table.add_row()
        _prevent_row_split(row)
        for column, (number, question) in enumerate(entries[offset : offset + 2]):
            cell = row.cells[column]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            labels = (
                [option.label for option in question.options]
                if question.options
                else ["正确", "错误"]
            )
            paragraph = cell.paragraphs[0]
            paragraph.add_run(f"{number}. ").bold = True
            paragraph.add_run("  ".join(f"□{label}" for label in labels))


def _answer_line(document: DocumentObject, prefix: str, line_count: int) -> None:
    for line in range(line_count):
        text = prefix if line == 0 else ""
        paragraph = document.add_paragraph(f"{text}{'_' * 78}")
        paragraph.paragraph_format.space_after = Pt(5)


def render_answer_sheet_document(
    request: AnswerSheetWordRequest,
    *,
    questions: Mapping[str, Question],
    template_path: Path,
) -> bytes:
    document = _open_template(template_path, "answer sheet")
    _add_identity(document, request)
    number_by_id = {
        assembled.question_id: number
        for number, assembled in ordered_assembled_questions(request.paper)
    }
    for section_position, assembled_section in enumerate(request.paper.sections):
        heading = document.add_paragraph(
            f"{_section_number(section_position)}、"
            f"{SECTION_TITLES[assembled_section.question_type]}",
            style="PaperSection",
        )
        heading.paragraph_format.keep_with_next = True
        entries: list[tuple[int, Question]] = []
        for assembled in assembled_section.questions:
            question = questions.get(assembled.question_id)
            if question is None:
                raise StudentPaperExportDataError(
                    f"answer sheet question is missing: {assembled.question_id}"
                )
            if question.question_type != assembled.question_type:
                raise StudentPaperExportDataError(
                    f"question type changed after assembly: {assembled.question_id}"
                )
            entries.append((number_by_id[assembled.question_id], question))

        if assembled_section.question_type in OBJECTIVE_TYPES:
            _add_objective_answer_rows(document, entries)
            continue
        for number, question in entries:
            if question.subquestions:
                label = document.add_paragraph(f"{number}. 综合题")
                label.paragraph_format.keep_with_next = True
                for child_index, _ in enumerate(question.subquestions, start=1):
                    _answer_line(document, f"（{child_index}）", 3)
            else:
                lines = 1 if question.question_type == QuestionType.FILL_BLANK else 5
                _answer_line(document, f"{number}. ", lines)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


async def _load_questions(
    session: AsyncSession,
    question_ids: list[str],
) -> dict[str, Question]:
    return {
        question_id: await recover_question(session, question_id)
        for question_id in question_ids
    }


async def _load_teacher_metadata(
    session: AsyncSession,
    question_ids: list[str],
) -> dict[str, TeacherQuestionMetadata]:
    knowledge_rows = (
        await session.execute(
            select(
                QuestionKnowledgePointModel.question_id,
                KnowledgePointModel.code,
                KnowledgePointModel.name,
            )
            .join(
                KnowledgePointModel,
                KnowledgePointModel.code
                == QuestionKnowledgePointModel.knowledge_point_code,
            )
            .where(
                QuestionKnowledgePointModel.question_id.in_(question_ids),
                KnowledgePointModel.is_active.is_(True),
            )
            .order_by(
                QuestionKnowledgePointModel.question_id,
                KnowledgePointModel.code,
            )
        )
    ).all()
    points: dict[str, list[tuple[str, str]]] = {}
    for question_id, code, name in knowledge_rows:
        points.setdefault(question_id, []).append((code, name))

    analyses = list(
        await session.scalars(
            select(QuestionAnalysisModel)
            .where(
                QuestionAnalysisModel.question_id.in_(question_ids),
                QuestionAnalysisModel.analysis_type == "difficulty_estimation",
            )
            .order_by(
                QuestionAnalysisModel.question_id,
                QuestionAnalysisModel.created_at.desc(),
                QuestionAnalysisModel.analysis_id.desc(),
            )
        )
    )
    latest: dict[str, DifficultyEstimationLLMOutput] = {}
    for analysis in analyses:
        if analysis.question_id in latest:
            continue
        try:
            latest[analysis.question_id] = DifficultyEstimationLLMOutput.model_validate(
                analysis.result
            )
        except ValidationError as exc:
            raise StudentPaperExportDataError(
                f"difficulty analysis is invalid: {analysis.analysis_id}"
            ) from exc

    metadata: dict[str, TeacherQuestionMetadata] = {}
    for question_id in question_ids:
        difficulty = latest.get(question_id)
        if difficulty is None:
            raise StudentPaperExportDataError(
                f"difficulty analysis is missing for question {question_id}"
            )
        metadata[question_id] = TeacherQuestionMetadata(
            knowledge_points=tuple(points.get(question_id, ())),
            difficulty_level=difficulty.difficulty_level,
            estimated_correct_rate=difficulty.estimated_correct_rate,
        )
    return metadata


async def export_teacher_answer_word(
    session: AsyncSession,
    request: TeacherAnswerWordRequest,
    *,
    storage_root: Path,
    template_root: Path,
) -> StudentPaperWordFile:
    question_ids = [
        assembled.question_id
        for _, assembled in ordered_assembled_questions(request.paper)
    ]
    questions = await _load_questions(session, question_ids)
    metadata = await _load_teacher_metadata(session, question_ids)
    image_data: dict[str, bytes] = {}
    for question in questions.values():
        for image in _question_images(question):
            if image.resource_id in image_data:
                continue
            stored = await get_stored_resource_file(
                session,
                image.resource_id,
                storage_root=storage_root,
            )
            image_data[image.resource_id] = await asyncio.to_thread(
                stored.path.read_bytes
            )
    content = await asyncio.to_thread(
        render_teacher_answer_document,
        request,
        questions=questions,
        metadata=metadata,
        image_data=image_data,
        template_path=template_root / TEACHER_TEMPLATE_FILENAME,
    )
    return StudentPaperWordFile(
        content=content,
        filename=safe_docx_filename(
            request.filename or request.title,
            "teacher-answer",
        ),
        allow_answers=True,
    )


async def export_answer_sheet_word(
    session: AsyncSession,
    request: AnswerSheetWordRequest,
    *,
    template_root: Path,
) -> StudentPaperWordFile:
    question_ids = [
        assembled.question_id
        for _, assembled in ordered_assembled_questions(request.paper)
    ]
    questions = await _load_questions(session, question_ids)
    content = await asyncio.to_thread(
        render_answer_sheet_document,
        request,
        questions=questions,
        template_path=template_root / ANSWER_SHEET_TEMPLATE_FILENAME,
    )
    return StudentPaperWordFile(
        content=content,
        filename=safe_docx_filename(
            request.filename or request.title,
            "answer-sheet",
        ),
        allow_answers=False,
        protected_texts=protected_answer_texts(questions),
    )


__all__ = [
    "ANSWER_SHEET_TEMPLATE_FILENAME",
    "DOCX_MIME_TYPE",
    "TEACHER_TEMPLATE_FILENAME",
    "TeacherQuestionMetadata",
    "create_answer_sheet_template",
    "create_teacher_answer_template",
    "export_answer_sheet_word",
    "export_teacher_answer_word",
    "render_answer_sheet_document",
    "render_teacher_answer_document",
]
