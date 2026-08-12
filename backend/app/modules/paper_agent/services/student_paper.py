"""A4 student-paper Word rendering with verified question images."""

from __future__ import annotations

import asyncio
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Mapping

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.paper_agent.schemas.assembly import (
    AssembledQuestion,
    PaperAssemblyResult,
)
from app.modules.paper_agent.schemas.question import Question, QuestionImage, QuestionType
from app.modules.paper_agent.schemas.student_paper import StudentPaperWordRequest
from app.modules.paper_agent.services.resource_storage import (
    get_stored_resource_file,
    recover_question,
)


DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
TEMPLATE_FILENAME = "student_paper.docx"
SECTION_TITLES = {
    QuestionType.SINGLE_CHOICE: "单项选择题",
    QuestionType.MULTIPLE_CHOICE: "多项选择题",
    QuestionType.TRUE_FALSE: "判断题",
    QuestionType.FILL_BLANK: "填空题",
    QuestionType.SHORT_ANSWER: "简答题",
    QuestionType.COMPOSITE: "综合题",
}
CHINESE_SECTION_NUMBERS = "一二三四五六七八九十"
INVALID_FILENAME_CHARACTERS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


class StudentPaperTemplateError(RuntimeError):
    pass


class StudentPaperExportDataError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class StudentPaperWordFile:
    content: bytes
    filename: str
    allow_answers: bool = True
    protected_texts: tuple[str, ...] = ()


def ordered_assembled_questions(
    paper: PaperAssemblyResult,
) -> list[tuple[int, AssembledQuestion]]:
    return [
        (number, question)
        for number, question in enumerate(
            (
                question
                for section in paper.sections
                for question in section.questions
            ),
            start=1,
        )
    ]


def safe_docx_filename(requested_name: str, fallback: str) -> str:
    safe_name = INVALID_FILENAME_CHARACTERS.sub("_", requested_name).rstrip(" .")
    return f"{safe_name or fallback}.docx"


def _set_cell_margins(cell, *, top=80, left=80, bottom=80, right=80) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in {
        "top": top,
        "left": left,
        "bottom": bottom,
        "right": right,
    }.items():
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def _prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def _set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _set_style_east_asia_font(style, name: str) -> None:
    style.font.name = "Times New Roman"
    run_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    run_fonts.set(qn("w:eastAsia"), name)


def _set_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field = OxmlElement("w:instrText")
    field.set(qn("xml:space"), "preserve")
    field.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, field, separate, text, end])


def _configure_a4_section(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    paragraph = section.footer.paragraphs[0]
    for child in list(paragraph._p):
        paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("第 ")
    _set_field(paragraph, "PAGE")
    paragraph.add_run(" 页，共 ")
    _set_field(paragraph, "NUMPAGES")
    paragraph.add_run(" 页")
    for run in paragraph.runs:
        _set_east_asia_font(run, "宋体")
        run.font.size = Pt(9)


def _request_field_updates(document: DocumentObject) -> None:
    settings = document.settings.element
    current = settings.find(qn("w:updateFields"))
    if current is None:
        current = OxmlElement("w:updateFields")
        settings.append(current)
    current.set(qn("w:val"), "true")


def _ensure_styles(document: DocumentObject) -> None:
    normal = document.styles["Normal"]
    _set_style_east_asia_font(normal, "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(0)

    styles = document.styles
    if "PaperTitle" not in styles:
        style = styles.add_style("PaperTitle", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_east_asia_font(style, "黑体")
        style.font.size = Pt(18)
        style.font.bold = True
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "PaperSubtitle" not in styles:
        style = styles.add_style("PaperSubtitle", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_east_asia_font(style, "宋体")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "PaperSection" not in styles:
        style = styles.add_style("PaperSection", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_east_asia_font(style, "黑体")
        style.font.size = Pt(12)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True


def create_student_paper_template() -> DocumentObject:
    document = Document()
    _ensure_styles(document)
    _configure_a4_section(document.sections[0])
    _request_field_updates(document)
    properties = document.core_properties
    properties.title = "FlowGate 高中生物学生卷模板"
    properties.subject = "A4 student paper template"
    properties.author = "FlowGate"
    properties.keywords = "FlowGate,高中生物,学生卷,A4"
    return document


def _clear_template_body(document: DocumentObject) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _image_buffer(data: bytes) -> tuple[BytesIO, int, int]:
    try:
        with Image.open(BytesIO(data)) as image:
            image.load()
            width, height = image.size
            converted = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            buffer = BytesIO()
            converted.save(buffer, format="PNG")
    except Exception as exc:
        raise StudentPaperExportDataError("question image cannot be decoded") from exc
    buffer.seek(0)
    return buffer, width, height


def _add_image(paragraph, image: QuestionImage, data: bytes) -> None:
    buffer, pixel_width, pixel_height = _image_buffer(data)
    maximum_width = 150.0
    maximum_height = 85.0
    natural_width = max(20.0, pixel_width * 25.4 / 120.0)
    natural_height = max(12.0, pixel_height * 25.4 / 120.0)
    scale = min(1.0, maximum_width / natural_width, maximum_height / natural_height)
    run = paragraph.add_run()
    inline_shape = run.add_picture(buffer, width=Mm(natural_width * scale))
    inline_shape._inline.docPr.set(
        "descr", image.alt_text or image.caption or "题目图片"
    )
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True


def _add_images(
    container,
    images: list[QuestionImage],
    image_data: Mapping[str, bytes],
) -> None:
    for image in images:
        data = image_data.get(image.resource_id)
        if data is None:
            raise StudentPaperExportDataError(
                f"question image data is missing: {image.resource_id}"
            )
        paragraph = container.add_paragraph()
        _add_image(paragraph, image, data)
        if image.caption:
            caption = container.add_paragraph(image.caption)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_with_next = True


def _add_options(
    container,
    question: Question,
    image_data: Mapping[str, bytes],
    *,
    keep_with_next: bool = False,
) -> None:
    if not question.options:
        return
    column_count = 1 if any(
        option.images or len(option.content or "") > 36
        for option in question.options
    ) else 2
    table = container.add_table(rows=1, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _remove_table_borders(table)
    row = table.rows[0]
    _prevent_row_split(row)
    for index, option in enumerate(question.options):
        cell = row.cells[index % column_count]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_margins(cell, top=30, bottom=30, left=80, right=80)
        paragraph = (
            cell.paragraphs[0]
            if index < column_count
            else cell.add_paragraph()
        )
        paragraph.paragraph_format.keep_with_next = bool(
            option.images
            or (keep_with_next and index == len(question.options) - 1)
        )
        label = paragraph.add_run(f"{option.label}. ")
        label.bold = True
        if option.content:
            paragraph.add_run(option.content)
        _add_images(cell, option.images, image_data)


def _add_answer_space(container, question_type: QuestionType) -> None:
    line_count = {
        QuestionType.FILL_BLANK: 1,
        QuestionType.SHORT_ANSWER: 4,
    }.get(question_type, 0)
    for _ in range(line_count):
        paragraph = container.add_paragraph("答：")
        paragraph.paragraph_format.space_after = Pt(8)


def _add_subquestions(
    container,
    question: Question,
    image_data: Mapping[str, bytes],
) -> None:
    for index, child in enumerate(question.subquestions, start=1):
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.keep_with_next = bool(child.images or child.options)
        number = paragraph.add_run(f"（{index}）")
        number.bold = True
        paragraph.add_run(child.stem)
        _add_images(container, child.images, image_data)
        _add_options(container, child, image_data)
        _add_answer_space(container, child.question_type)


def _add_question_block(
    document: DocumentObject,
    *,
    number: int,
    score: int,
    question: Question,
    image_data: Mapping[str, bytes],
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.widow_control = True
    paragraph.paragraph_format.keep_with_next = bool(
        question.images or question.options or question.subquestions
    )
    number_run = paragraph.add_run(f"{number}. ")
    number_run.bold = True
    paragraph.add_run(question.stem)
    score_run = paragraph.add_run(f"（{score} 分）")
    score_run.font.size = Pt(9)
    _add_images(document, question.images, image_data)
    _add_options(document, question, image_data)
    _add_subquestions(document, question, image_data)
    _add_answer_space(document, question.question_type)


def _section_number(index: int) -> str:
    if index < len(CHINESE_SECTION_NUMBERS):
        return CHINESE_SECTION_NUMBERS[index]
    return str(index + 1)


def _add_cover(document: DocumentObject, request: StudentPaperWordRequest) -> None:
    title = document.add_paragraph(request.title, style="PaperTitle")
    title.paragraph_format.keep_with_next = True
    subtitle_text = request.subtitle or f"{request.paper.module_code} 学生卷"
    subtitle = document.add_paragraph(subtitle_text, style="PaperSubtitle")
    subtitle.paragraph_format.keep_with_next = True

    identity = document.add_table(rows=2, cols=2)
    identity.alignment = WD_TABLE_ALIGNMENT.CENTER
    identity.autofit = True
    _remove_table_borders(identity)
    values = [
        f"学校：{request.school_name or '____________'}",
        f"年级：{request.grade_name}",
        "姓名：____________",
        "考号：____________",
    ]
    cells = [cell for row in identity.rows for cell in row.cells]
    for cell, value in zip(cells, values, strict=True):
        _set_cell_margins(cell, top=30, bottom=30, left=60, right=60)
        cell.paragraphs[0].add_run(value)

    overview = document.add_paragraph()
    overview.alignment = WD_ALIGN_PARAGRAPH.CENTER
    overview.add_run(
        f"考试时间：{request.duration_minutes} 分钟    "
        f"满分：{request.paper.total_score} 分    "
        f"共 {request.paper.question_count} 题"
    ).bold = True
    overview.paragraph_format.space_after = Pt(6)

    if request.instructions:
        notice = document.add_table(rows=1, cols=1)
        notice.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = notice.cell(0, 0)
        _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.paragraphs[0].add_run("注意事项").bold = True
        for index, instruction in enumerate(request.instructions, start=1):
            cell.add_paragraph(f"{index}. {instruction}")


def render_student_paper_document(
    request: StudentPaperWordRequest,
    *,
    questions: Mapping[str, Question],
    image_data: Mapping[str, bytes],
    template_path: Path,
) -> bytes:
    if not template_path.is_file():
        raise StudentPaperTemplateError(
            f"student paper template is missing: {template_path}"
        )
    try:
        document = Document(template_path)
    except Exception as exc:
        raise StudentPaperTemplateError("student paper template cannot be opened") from exc
    _clear_template_body(document)
    _ensure_styles(document)
    for section in document.sections:
        _configure_a4_section(section)
    _request_field_updates(document)
    _add_cover(document, request)

    number_by_id = {
        assembled.question_id: number
        for number, assembled in ordered_assembled_questions(request.paper)
    }
    for section_position, assembled_section in enumerate(request.paper.sections):
        if section_position > 0 and request.start_each_section_on_new_page:
            document.add_page_break()
        title = SECTION_TITLES[assembled_section.question_type]
        heading = document.add_paragraph(
            f"{_section_number(section_position)}、{title}",
            style="PaperSection",
        )
        heading.paragraph_format.keep_with_next = True
        description = document.add_paragraph(
            f"本大题共 {assembled_section.count} 题，每题 "
            f"{assembled_section.score_per_question} 分，共 "
            f"{assembled_section.count * assembled_section.score_per_question} 分。"
        )
        description.paragraph_format.keep_with_next = True
        for assembled_question in assembled_section.questions:
            question = questions.get(assembled_question.question_id)
            if question is None:
                raise StudentPaperExportDataError(
                    f"assembled question is missing: {assembled_question.question_id}"
                )
            if question.question_type != assembled_question.question_type:
                raise StudentPaperExportDataError(
                    f"question type changed after assembly: {assembled_question.question_id}"
                )
            _add_question_block(
                document,
                number=number_by_id[assembled_question.question_id],
                score=assembled_question.score,
                question=question,
                image_data=image_data,
            )

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _question_images(question: Question) -> list[QuestionImage]:
    images = list(question.images)
    for option in question.options:
        images.extend(option.images)
    for child in question.subquestions:
        images.extend(_question_images(child))
    return images


def protected_answer_texts(questions: Mapping[str, Question]) -> tuple[str, ...]:
    values: list[str] = []
    public_texts: list[str] = []

    def collect(question: Question) -> None:
        public_texts.append(question.stem)
        public_texts.extend(option.content or "" for option in question.options)
        if isinstance(question.answer, list):
            values.extend(question.answer)
        elif question.answer:
            values.append(question.answer)
        if question.explanation:
            values.append(question.explanation)
        for child in question.subquestions:
            collect(child)

    for question in questions.values():
        collect(question)
    normalized_public = re.sub(r"\s+", "", "".join(public_texts))
    protected = [
        value
        for value in values
        if len(re.sub(r"\s+", "", value)) >= 8
        and re.sub(r"\s+", "", value) not in normalized_public
    ]
    return tuple(dict.fromkeys(protected))


async def export_student_paper_word(
    session: AsyncSession,
    request: StudentPaperWordRequest,
    *,
    storage_root: Path,
    template_root: Path,
) -> StudentPaperWordFile:
    assembled_questions: list[AssembledQuestion] = [
        question
        for section in request.paper.sections
        for question in section.questions
    ]
    questions: dict[str, Question] = {}
    for assembled in assembled_questions:
        questions[assembled.question_id] = await recover_question(
            session, assembled.question_id
        )

    image_data: dict[str, bytes] = {}
    for question in questions.values():
        for image in _question_images(question):
            if image.resource_id in image_data:
                continue
            stored = await get_stored_resource_file(
                session,
                image.resource_id,
                storage_root=storage_root,
            )
            image_data[image.resource_id] = await asyncio.to_thread(stored.path.read_bytes)

    content = await asyncio.to_thread(
        render_student_paper_document,
        request,
        questions=questions,
        image_data=image_data,
        template_path=template_root / TEMPLATE_FILENAME,
    )
    return StudentPaperWordFile(
        content=content,
        filename=safe_docx_filename(
            request.filename or request.title,
            "student-paper",
        ),
        allow_answers=False,
        protected_texts=protected_answer_texts(questions),
    )
