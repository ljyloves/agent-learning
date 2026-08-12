import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image
from pydantic import ValidationError

from app.modules.paper_agent.schemas.assembly import (
    AssembledQuestion,
    AssembledSection,
    PaperAssemblyResult,
)
from app.modules.paper_agent.schemas.question import (
    Question,
    QuestionDifficulty,
    QuestionImage,
    QuestionOption,
    QuestionType,
)
from app.modules.paper_agent.schemas.source import QuestionSource, SourceType
from app.modules.paper_agent.schemas.student_paper import StudentPaperWordRequest
from app.modules.paper_agent.services.student_paper import (
    StudentPaperExportDataError,
    StudentPaperTemplateError,
    create_student_paper_template,
    protected_answer_texts,
    render_student_paper_document,
)


def source() -> QuestionSource:
    return QuestionSource(
        source_id="source-word-export",
        source_type=SourceType.MANUAL,
        name="BIO-031 test source",
        external_id="BIO-031",
    )


def image_bytes() -> bytes:
    output = BytesIO()
    Image.new("RGB", (1200, 480), color=(245, 250, 255)).save(output, "PNG")
    return output.getvalue()


def paper() -> PaperAssemblyResult:
    easy_choice = AssembledQuestion(
        question_id="question-1",
        question_type=QuestionType.SINGLE_CHOICE,
        difficulty=QuestionDifficulty.EASY,
        score=3,
    )
    medium_choice = AssembledQuestion(
        question_id="question-2",
        question_type=QuestionType.SINGLE_CHOICE,
        difficulty=QuestionDifficulty.MEDIUM,
        score=4,
    )
    fill_blank = AssembledQuestion(
        question_id="question-3",
        question_type=QuestionType.FILL_BLANK,
        difficulty=QuestionDifficulty.MEDIUM,
        score=5,
    )
    return PaperAssemblyResult(
        module_code="BIO-M1",
        question_count=3,
        total_score=12,
        sections=[
            AssembledSection(
                section_index=0,
                question_type=QuestionType.SINGLE_CHOICE,
                difficulty=QuestionDifficulty.EASY,
                count=1,
                score_per_question=3,
                questions=[easy_choice],
            ),
            AssembledSection(
                section_index=1,
                question_type=QuestionType.SINGLE_CHOICE,
                difficulty=QuestionDifficulty.MEDIUM,
                count=1,
                score_per_question=4,
                questions=[medium_choice],
            ),
            AssembledSection(
                section_index=2,
                question_type=QuestionType.FILL_BLANK,
                difficulty=QuestionDifficulty.MEDIUM,
                count=1,
                score_per_question=5,
                questions=[fill_blank],
            ),
        ],
    )


def questions() -> dict[str, Question]:
    provenance = source()
    question_image = QuestionImage(
        resource_id="image-1",
        uri="storage://questions/image-1.png",
        source=provenance,
        mime_type="image/png",
        alt_text="细胞结构示意图",
        caption="图 1 细胞结构示意图",
    )
    base_options = [
        QuestionOption(label="A", content="细胞膜"),
        QuestionOption(label="B", content="细胞壁"),
        QuestionOption(label="C", content="细胞核"),
        QuestionOption(label="D", content="细胞质"),
    ]
    return {
        "question-1": Question(
            id="question-1",
            question_type=QuestionType.SINGLE_CHOICE,
            difficulty=QuestionDifficulty.EASY,
            stem="观察图示，控制物质进出细胞的结构是",
            source=provenance,
            options=base_options,
            answer="A-绝密答案",
            explanation="绝密解析不应出现在学生卷中",
            images=[question_image],
        ),
        "question-2": Question(
            id="question-2",
            question_type=QuestionType.SINGLE_CHOICE,
            difficulty=QuestionDifficulty.MEDIUM,
            stem="下列关于细胞结构与功能的叙述，正确的是",
            source=provenance,
            options=[
                QuestionOption(
                    label="A",
                    content="这是一条超过阈值的长选项，用于验证长文本自动改为单列排版，避免在窄单元格内拥挤。",
                ),
                QuestionOption(label="B", content="核糖体参与蛋白质合成"),
            ],
            answer="B",
            explanation="绝密解析",
        ),
        "question-3": Question(
            id="question-3",
            question_type=QuestionType.FILL_BLANK,
            difficulty=QuestionDifficulty.MEDIUM,
            stem="遗传信息主要储存在细胞的________中。",
            source=provenance,
            answer="细胞核",
            explanation="绝密解析",
        ),
    }


class StudentPaperWordTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.template_path = Path(self.temporary.name) / "student_paper.docx"
        create_student_paper_template().save(self.template_path)

    def tearDown(self):
        self.temporary.cleanup()

    def test_answer_leak_fingerprints_exclude_public_question_content(self):
        question_map = questions()
        first = question_map["question-1"]
        public_explanation = first.stem
        question_map["question-1"] = first.model_copy(
            update={"explanation": public_explanation}
        )
        second = question_map["question-2"]
        private_explanation = "这是一段只允许教师查看的详细解析内容"
        question_map["question-2"] = second.model_copy(
            update={"explanation": private_explanation}
        )

        protected = protected_answer_texts(question_map)

        self.assertNotIn(public_explanation, protected)
        self.assertIn(private_explanation, protected)

    def test_renders_a4_numbered_images_and_page_controls(self):
        request = StudentPaperWordRequest(
            paper=paper(),
            title="高一生物阶段测试",
            duration_minutes=45,
        )
        content = render_student_paper_document(
            request,
            questions=questions(),
            image_data={"image-1": image_bytes()},
            template_path=self.template_path,
        )
        document = Document(BytesIO(content))
        option_tables = [
            table
            for table in document.tables
            if any(
                "A. " in paragraph.text
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            )
        ]
        self.assertTrue(option_tables)
        self.assertTrue(all(len(table.rows) == 1 for table in option_tables))
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, delta=0.2)
        self.assertAlmostEqual(section.page_height.mm, 297, delta=0.2)
        self.assertAlmostEqual(section.left_margin.mm, 18, delta=0.2)
        self.assertAlmostEqual(section.right_margin.mm, 18, delta=0.2)

        xml = document._element.xml
        paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("1. 观察图示", paragraphs)
        self.assertIn("2. 下列关于", paragraphs)
        self.assertIn("3. 遗传信息", paragraphs)
        self.assertNotIn("绝密答案", xml)
        self.assertNotIn("绝密解析", xml)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertLessEqual(document.inline_shapes[0].width, Mm(150))
        self.assertIn('descr="细胞结构示意图"', xml)

        page_breaks = document._element.xpath('.//w:br[@w:type="page"]')
        self.assertEqual(len(page_breaks), 2)
        self.assertTrue(document._element.xpath(".//w:cantSplit"))
        self.assertTrue(document._element.xpath(".//w:keepNext"))
        self.assertTrue(document._element.xpath(".//w:tbl[w:tblGrid[count(w:gridCol)=1]]"))

        footer_xml = section.footer._element.xml
        self.assertIn("PAGE", footer_xml)
        self.assertIn("NUMPAGES", footer_xml)
        update_fields = document.settings.element.find(qn("w:updateFields"))
        self.assertIsNotNone(update_fields)
        self.assertEqual(update_fields.get(qn("w:val")), "true")

    def test_rejects_missing_template_and_image_data(self):
        request = StudentPaperWordRequest(paper=paper())
        with self.assertRaises(StudentPaperTemplateError):
            render_student_paper_document(
                request,
                questions=questions(),
                image_data={"image-1": image_bytes()},
                template_path=Path(self.temporary.name) / "missing.docx",
            )
        with self.assertRaises(StudentPaperExportDataError):
            render_student_paper_document(
                request,
                questions=questions(),
                image_data={},
                template_path=self.template_path,
            )

    def test_filename_contract_rejects_unsafe_name(self):
        with self.assertRaises(ValidationError):
            StudentPaperWordRequest(paper=paper(), filename="biology/test.docx")
        request = StudentPaperWordRequest(paper=paper(), filename="student.docx")
        self.assertEqual(request.filename, "student")


if __name__ == "__main__":
    unittest.main()
