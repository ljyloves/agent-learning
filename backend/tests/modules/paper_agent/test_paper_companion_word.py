import re
import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml.ns import qn

from app.modules.paper_agent.schemas.assembly import (
    AssembledQuestion,
    AssembledSection,
    PaperAssemblyResult,
)
from app.modules.paper_agent.schemas.question import (
    Question,
    QuestionDifficulty,
    QuestionOption,
    QuestionType,
)
from app.modules.paper_agent.schemas.source import QuestionSource, SourceType
from app.modules.paper_agent.schemas.student_paper import (
    AnswerSheetWordRequest,
    StudentPaperWordRequest,
    TeacherAnswerWordRequest,
)
from app.modules.paper_agent.services.paper_companions import (
    TeacherQuestionMetadata,
    create_answer_sheet_template,
    create_teacher_answer_template,
    render_answer_sheet_document,
    render_teacher_answer_document,
)
from app.modules.paper_agent.services.student_paper import (
    StudentPaperExportDataError,
    create_student_paper_template,
    ordered_assembled_questions,
    render_student_paper_document,
)


def paper() -> PaperAssemblyResult:
    return PaperAssemblyResult(
        module_code="BIO-M1",
        question_count=3,
        total_score=15,
        sections=[
            AssembledSection(
                section_index=0,
                question_type=QuestionType.SINGLE_CHOICE,
                difficulty=QuestionDifficulty.MEDIUM,
                count=2,
                score_per_question=4,
                questions=[
                    AssembledQuestion(
                        question_id=f"q-{index}",
                        question_type=QuestionType.SINGLE_CHOICE,
                        difficulty=QuestionDifficulty.MEDIUM,
                        score=4,
                    )
                    for index in (1, 2)
                ],
            ),
            AssembledSection(
                section_index=1,
                question_type=QuestionType.SHORT_ANSWER,
                difficulty=QuestionDifficulty.HARD,
                count=1,
                score_per_question=7,
                questions=[
                    AssembledQuestion(
                        question_id="q-3",
                        question_type=QuestionType.SHORT_ANSWER,
                        difficulty=QuestionDifficulty.HARD,
                        score=7,
                    )
                ],
            ),
        ],
    )


def questions() -> dict[str, Question]:
    source = QuestionSource(
        source_id="source-companion",
        source_type=SourceType.MANUAL,
        name="companion test source",
        external_id="BIO-032-033",
    )
    options = [
        QuestionOption(label="A", content="细胞膜"),
        QuestionOption(label="B", content="细胞核"),
        QuestionOption(label="C", content="核糖体"),
        QuestionOption(label="D", content="线粒体"),
    ]
    return {
        "q-1": Question(
            id="q-1",
            question_type=QuestionType.SINGLE_CHOICE,
            difficulty=QuestionDifficulty.MEDIUM,
            stem="控制物质进出细胞的结构是",
            source=source,
            options=options,
            answer="A",
            explanation="细胞膜具有选择透过性。",
        ),
        "q-2": Question(
            id="q-2",
            question_type=QuestionType.SINGLE_CHOICE,
            difficulty=QuestionDifficulty.MEDIUM,
            stem="蛋白质合成的场所是",
            source=source,
            options=options,
            answer="C",
            explanation="核糖体是蛋白质合成场所。",
        ),
        "q-3": Question(
            id="q-3",
            question_type=QuestionType.SHORT_ANSWER,
            difficulty=QuestionDifficulty.HARD,
            stem="说明细胞膜的结构特点。",
            source=source,
            answer="具有一定的流动性。",
            explanation="磷脂和多数蛋白质可以运动。",
        ),
    }


def metadata() -> dict[str, TeacherQuestionMetadata]:
    return {
        "q-1": TeacherQuestionMetadata((("BIO-M1-K02", "细胞膜"),), 2, 0.72),
        "q-2": TeacherQuestionMetadata((("BIO-M1-K03", "蛋白质"),), 3, 0.55),
        "q-3": TeacherQuestionMetadata((("BIO-M1-K02", "细胞膜"),), 4, 0.30),
    }


def all_text(document) -> str:
    values = []
    for paragraph in document._element.iter(qn("w:p")):
        values.append(
            "".join(node.text or "" for node in paragraph.iter(qn("w:t")))
        )
    return "\n".join(values)


class PaperCompanionWordTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        root = Path(self.temporary.name)
        self.student_template = root / "student_paper.docx"
        self.teacher_template = root / "teacher_answer.docx"
        self.answer_sheet_template = root / "answer_sheet.docx"
        create_student_paper_template().save(self.student_template)
        create_teacher_answer_template().save(self.teacher_template)
        create_answer_sheet_template().save(self.answer_sheet_template)

    def tearDown(self):
        self.temporary.cleanup()

    def test_teacher_document_contains_complete_marking_context(self):
        content = render_teacher_answer_document(
            TeacherAnswerWordRequest(paper=paper()),
            questions=questions(),
            metadata=metadata(),
            image_data={},
            template_path=self.teacher_template,
        )
        document = Document(BytesIO(content))
        text = all_text(document)
        self.assertIn("答案：A", text)
        self.assertIn("解析：细胞膜具有选择透过性。", text)
        self.assertIn("知识点：BIO-M1-K02 细胞膜", text)
        self.assertIn("难度：2 级（预计正确率 72%）", text)
        self.assertIn("3. 说明细胞膜的结构特点。", text)
        teacher_detail_tables = [
            table
            for table in document.tables
            if "答案：" in all_text(table)
        ]
        self.assertEqual(len(teacher_detail_tables), 3)
        self.assertTrue(
            all(len(table.rows) == 1 for table in teacher_detail_tables)
        )
        self.assertAlmostEqual(document.sections[0].page_width.mm, 210, delta=0.2)

    def test_teacher_document_rejects_incomplete_required_fields(self):
        incomplete = questions()
        incomplete["q-2"] = incomplete["q-2"].model_copy(
            update={"explanation": None}
        )
        with self.assertRaises(StudentPaperExportDataError):
            render_teacher_answer_document(
                TeacherAnswerWordRequest(paper=paper()),
                questions=incomplete,
                metadata=metadata(),
                image_data={},
                template_path=self.teacher_template,
            )
        missing_points = metadata()
        missing_points["q-2"] = TeacherQuestionMetadata((), 3, 0.55)
        with self.assertRaises(StudentPaperExportDataError):
            render_teacher_answer_document(
                TeacherAnswerWordRequest(paper=paper()),
                questions=questions(),
                metadata=missing_points,
                image_data={},
                template_path=self.teacher_template,
            )

    def test_answer_sheet_numbers_exactly_match_student_paper(self):
        student = Document(
            BytesIO(
                render_student_paper_document(
                    StudentPaperWordRequest(paper=paper()),
                    questions=questions(),
                    image_data={},
                    template_path=self.student_template,
                )
            )
        )
        sheet = Document(
            BytesIO(
                render_answer_sheet_document(
                    AnswerSheetWordRequest(paper=paper()),
                    questions=questions(),
                    template_path=self.answer_sheet_template,
                )
            )
        )
        expected = [number for number, _ in ordered_assembled_questions(paper())]
        student_numbers = [
            int(match.group(1))
            for match in map(
                lambda paragraph: re.match(r"^(\d+)\. ", paragraph.text),
                student.paragraphs,
            )
            if match
        ]
        sheet_numbers = sorted(
            {
                int(match.group(1))
                for match in re.finditer(r"(?m)^(\d+)\. ", all_text(sheet))
            }
        )
        self.assertEqual(student_numbers, expected)
        self.assertEqual(sheet_numbers, expected)
        self.assertIn("□A", all_text(sheet))
        self.assertNotIn("细胞膜具有选择透过性", all_text(sheet))
        self.assertAlmostEqual(sheet.sections[0].page_height.mm, 297, delta=0.2)


if __name__ == "__main__":
    unittest.main()
