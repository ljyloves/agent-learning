import shutil
import unittest
from io import BytesIO

from docx import Document
from docx.shared import Mm
from PIL import Image
from pypdf import PdfReader, PdfWriter

from app.modules.paper_agent.services.document_conversion import (
    DocumentConversionError,
    _assert_content_consistency,
    _comparison_text,
    convert_word_export_to_pdf,
)
from app.modules.paper_agent.services.render_validation import (
    DocumentRenderValidationError,
    inspect_rendered_document,
)
from app.modules.paper_agent.services.student_paper import StudentPaperWordFile


def word_file(
    text: str,
    filename: str = "biology.docx",
    *,
    include_teacher_detail: bool = True,
    include_image: bool = False,
    allow_answers: bool = True,
) -> StudentPaperWordFile:
    document = Document()
    document.add_heading("高中生物测试", level=1)
    document.add_paragraph(text)
    if include_image:
        image_data = BytesIO()
        Image.new("RGB", (80, 60), "white").save(image_data, format="PNG")
        image_data.seek(0)
        document.add_picture(image_data, width=Mm(20))
    if include_teacher_detail:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "答案：A；知识点：细胞膜；难度：2级"
    output = BytesIO()
    document.save(output)
    return StudentPaperWordFile(
        content=output.getvalue(),
        filename=filename,
        allow_answers=allow_answers,
    )


@unittest.skipUnless(
    shutil.which("soffice") or shutil.which("libreoffice"),
    "LibreOffice is not installed in this test environment",
)
class DocumentConversionTests(unittest.IsolatedAsyncioTestCase):
    def test_content_comparison_ignores_pdf_punctuation_reordering(self):
        docx_text = "1. 某同学复习“线粒体功能”时整理了四条笔记。"
        pdf_text = "1. “”某同学复习线粒体功能时整理了四条笔记。"

        self.assertEqual(_comparison_text(docx_text), _comparison_text(pdf_text))

    async def test_converts_docx_to_downloadable_text_equivalent_pdf(self):
        source = word_file("1. 控制物质进出细胞的结构是细胞膜。")
        result = await convert_word_export_to_pdf(source)

        self.assertTrue(result.content.startswith(b"%PDF-"))
        self.assertEqual(result.filename, "biology.pdf")
        self.assertGreaterEqual(result.page_count, 1)
        self.assertTrue(result.inspection.passed)
        self.assertEqual(result.inspection.blank_pages, ())
        self.assertEqual(result.inspection.answer_leaks, ())
        self.assertEqual(result.inspection.abnormal_break_pages, ())
        reader = PdfReader(BytesIO(result.content))
        text = "".join(page.extract_text() or "" for page in reader.pages)
        self.assertIn("高中生物测试", text)
        self.assertIn("细胞膜", text)
        self.assertIn("答案", text)

    async def test_content_validator_rejects_mismatched_documents(self):
        source = word_file("原始题目内容")
        converted = await convert_word_export_to_pdf(source)
        different = word_file("另一份完全不同的题目内容")

        with self.assertRaises(DocumentConversionError):
            _assert_content_consistency(different.content, converted.content)

    async def test_accepts_all_images_rendered(self):
        source = word_file(
            "1. 观察细胞结构图并回答问题。",
            include_teacher_detail=False,
            include_image=True,
            allow_answers=False,
        )

        result = await convert_word_export_to_pdf(source)

        self.assertEqual(result.inspection.expected_image_count, 1)
        self.assertGreaterEqual(result.inspection.rendered_image_count, 1)

    async def test_rejects_answer_leak_in_student_document(self):
        source = word_file(
            "1. 控制物质进出细胞的结构是细胞膜。",
            allow_answers=False,
        )

        with self.assertRaisesRegex(DocumentConversionError, "answer leakage"):
            await convert_word_export_to_pdf(source)

    async def test_rejects_blank_page(self):
        source = word_file(
            "1. 控制物质进出细胞的结构是细胞膜。",
            include_teacher_detail=False,
            allow_answers=False,
        )
        converted = await convert_word_export_to_pdf(source)
        reader = PdfReader(BytesIO(converted.content))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.add_blank_page(width=595, height=842)
        output = BytesIO()
        writer.write(output)

        with self.assertRaisesRegex(
            DocumentRenderValidationError,
            "blank pages",
        ):
            inspect_rendered_document(
                source.content,
                output.getvalue(),
                allow_answers=False,
            )

    async def test_rejects_missing_rendered_image(self):
        with_image = word_file(
            "1. 观察细胞结构图并回答问题。",
            include_teacher_detail=False,
            include_image=True,
            allow_answers=False,
        )
        without_image = word_file(
            "1. 观察细胞结构图并回答问题。",
            include_teacher_detail=False,
            allow_answers=False,
        )
        converted = await convert_word_export_to_pdf(without_image)

        with self.assertRaisesRegex(
            DocumentRenderValidationError,
            "missing rendered images",
        ):
            inspect_rendered_document(
                with_image.content,
                converted.content,
                allow_answers=False,
            )

    async def test_rejects_orphan_section_heading_at_page_end(self):
        source = word_file(
            "二、简答题",
            include_teacher_detail=False,
            allow_answers=False,
        )

        with self.assertRaisesRegex(DocumentConversionError, "abnormal page breaks"):
            await convert_word_export_to_pdf(source)


if __name__ == "__main__":
    unittest.main()
